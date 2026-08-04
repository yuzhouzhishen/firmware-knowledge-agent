from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from pathlib import Path

from pydantic import BaseModel

from firmware_knowledge_agent.models import SourceSpec
from firmware_knowledge_agent.public_ingestion import (
    extract_document_markdown,
)


MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
SUPPORTED_SUFFIXES = {
    ".docx",
    ".htm",
    ".html",
    ".markdown",
    ".md",
    ".pdf",
    ".txt",
}
_SOURCE_ID_PATTERN = re.compile(r"^[a-z0-9][a-z0-9-]{1,63}$")
_CONTROL_CHARACTERS = re.compile(
    r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]"
)
_REDACTIONS = (
    (
        re.compile(
            r"https?://[^\s)\]>'\"]*"
            r"(?:corp\.ifanr\.com|ifanrprod\.com|thecandysign\.com)"
            r"[^\s)\]>'\"]*",
            re.IGNORECASE,
        ),
        "<REDACTED_INTERNAL_URL>",
    ),
    (
        re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b"),
        "<REDACTED_IP>",
    ),
    (
        re.compile(r"\b(?:[0-9A-F]{2}:){5}[0-9A-F]{2}\b", re.IGNORECASE),
        "<REDACTED_MAC>",
    ),
    (
        re.compile(
            r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"
        ),
        "<REDACTED_EMAIL>",
    ),
    (
        re.compile(
            r"\beyJ[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+\b"
        ),
        "<REDACTED_TOKEN>",
    ),
    (
        re.compile(r"\b\d{14,20}\b"),
        "<REDACTED_DEVICE_ID>",
    ),
    (
        re.compile(
            r"((?:api[_ -]?key|access[_ -]?token|secret|password|passwd|"
            r"psn|serial(?:[_ -]?number)?|device[_ -]?id|sn)"
            r"\s*[:=]\s*)[^\s,;}\"']+",
            re.IGNORECASE,
        ),
        r"\1<REDACTED_VALUE>",
    ),
)


class LocalIngestionResult(BaseModel):
    source_id: str
    title: str
    document_path: str
    catalog_path: str
    version: str
    content_characters: int
    created: bool


def ingest_document_bytes(
    *,
    filename: str,
    content: bytes,
    catalog_path: Path,
    source_id: str | None = None,
    title: str | None = None,
    source_url: str | None = None,
    component: str = "general",
    confidentiality: str = "private-local",
) -> LocalIngestionResult:
    if not filename.strip():
        raise ValueError("filename must not be empty")
    if not content:
        raise ValueError("uploaded document is empty")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("uploaded document exceeds the 10 MiB limit")

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise ValueError(f"unsupported document type; expected one of: {supported}")

    normalized_title = (title or Path(filename).stem).strip()
    if not normalized_title:
        raise ValueError("title must not be empty")
    normalized_id = _normalize_source_id(
        source_id,
        filename=filename,
        content=content,
    )
    markdown = sanitize_private_text(
        _extract_markdown(
            suffix,
            content,
            title=normalized_title,
        )
    )
    if len(markdown) < 100:
        raise ValueError("extracted document is too short to index")
    if not markdown.lstrip().startswith("#"):
        markdown = f"# {normalized_title}\n\n{markdown}"

    catalog_path = catalog_path.resolve()
    corpus_root = catalog_path.parent
    uploads_dir = corpus_root / "uploads"
    uploads_dir.mkdir(parents=True, exist_ok=True)
    relative_document_path = Path("uploads") / f"{normalized_id}.md"
    document_path = corpus_root / relative_document_path

    digest = hashlib.sha256(content).hexdigest()
    source = SourceSpec(
        id=normalized_id,
        title=normalized_title,
        path=relative_document_path.as_posix(),
        source_url=(
            source_url.strip()
            if source_url and source_url.strip()
            else f"local://{Path(filename).name}"
        ),
        version=f"sha256:{digest}",
        license_note=(
            "User-provided local document; private index only; "
            "do not redistribute."
        ),
        source_type=suffix.removeprefix("."),
        component=component.strip() or "general",
        confidentiality=confidentiality.strip() or "private-local",
    )
    catalog = _load_catalog(catalog_path)
    existing_ids = {item.id for item in catalog}
    updated_catalog = [
        item for item in catalog if item.id != source.id
    ] + [source]

    _atomic_write_text(document_path, f"{markdown.rstrip()}\n")
    _atomic_write_text(
        catalog_path,
        f"{json.dumps(
            [item.model_dump(mode='json') for item in updated_catalog],
            ensure_ascii=False,
            indent=2,
        )}\n",
    )
    return LocalIngestionResult(
        source_id=source.id,
        title=source.title,
        document_path=str(document_path),
        catalog_path=str(catalog_path),
        version=source.version,
        content_characters=len(markdown),
        created=source.id not in existing_ids,
    )


def sanitize_private_text(value: str) -> str:
    cleaned = _CONTROL_CHARACTERS.sub("", value)
    cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
    for pattern, replacement in _REDACTIONS:
        cleaned = pattern.sub(replacement, cleaned)
    cleaned = re.sub(r"[ \t]+\n", "\n", cleaned)
    return re.sub(r"\n{3,}", "\n\n", cleaned).strip()


def _extract_markdown(
    suffix: str,
    content: bytes,
    *,
    title: str,
) -> str:
    if suffix in {".md", ".markdown", ".txt"}:
        return content.decode("utf-8-sig")
    if suffix in {".htm", ".html"}:
        return extract_document_markdown(content.decode("utf-8-sig"))
    if suffix == ".pdf":
        return _extract_pdf(content, title=title)
    if suffix == ".docx":
        return _extract_docx(content, title=title)
    raise ValueError(f"unsupported document type: {suffix}")


def _extract_pdf(content: bytes, *, title: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages = [
        f"## Page {index}\n\n{text.strip()}"
        for index, page in enumerate(reader.pages, start=1)
        if (text := page.extract_text() or "").strip()
    ]
    return f"# {title}\n\n" + "\n\n".join(pages)


def _extract_docx(content: bytes, *, title: str) -> str:
    from docx import Document

    document = Document(io.BytesIO(content))
    lines = [f"# {title}"]
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name.casefold()
        if style_name.startswith("heading"):
            match = re.search(r"(\d+)", style_name)
            level = min(int(match.group(1)) if match else 2, 4)
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)
    for table in document.tables:
        rows = [
            [cell.text.strip().replace("|", "\\|") for cell in row.cells]
            for row in table.rows
        ]
        if not rows:
            continue
        lines.append("| " + " | ".join(rows[0]) + " |")
        lines.append("| " + " | ".join("---" for _ in rows[0]) + " |")
        lines.extend(
            "| " + " | ".join(row) + " |"
            for row in rows[1:]
        )
    return "\n\n".join(lines)


def _normalize_source_id(
    source_id: str | None,
    *,
    filename: str,
    content: bytes,
) -> str:
    if source_id is not None:
        normalized = source_id.strip().lower()
        if not _SOURCE_ID_PATTERN.fullmatch(normalized):
            raise ValueError(
                "source_id must use 2-64 lowercase letters, digits, or hyphens"
            )
        return normalized

    stem = Path(filename).stem.lower()
    slug = re.sub(r"[^a-z0-9]+", "-", stem).strip("-")[:48]
    if len(slug) >= 2:
        return slug
    digest = hashlib.sha256(content).hexdigest()[:12]
    return f"local-{digest}"


def _load_catalog(path: Path) -> list[SourceSpec]:
    if not path.exists():
        return []
    raw = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(raw, list):
        raise ValueError("catalog must contain a JSON array")
    return [SourceSpec.model_validate(item) for item in raw]


def _atomic_write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_suffix(f"{path.suffix}.tmp")
    temporary_path.write_text(content, encoding="utf-8")
    temporary_path.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Ingest a local document into a private RAG corpus.",
    )
    parser.add_argument("file", type=Path)
    parser.add_argument(
        "--catalog",
        type=Path,
        default=Path("var/private-corpus/sources.json"),
    )
    parser.add_argument("--source-id")
    parser.add_argument("--title")
    parser.add_argument("--source-url")
    parser.add_argument("--component", default="general")
    parser.add_argument("--confidentiality", default="private-local")
    args = parser.parse_args()
    result = ingest_document_bytes(
        filename=args.file.name,
        content=args.file.read_bytes(),
        catalog_path=args.catalog,
        source_id=args.source_id,
        title=args.title,
        source_url=args.source_url,
        component=args.component,
        confidentiality=args.confidentiality,
    )
    print(result.model_dump_json(indent=2))


if __name__ == "__main__":
    main()
